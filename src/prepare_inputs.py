#!/usr/bin/env python
"""Create the text-free input files that src/analysis.py needs.

Run this once on your own computer, where the news text is available. The Urdu News
Dataset 1M may only be used for non-commercial research with credit to the news sources,
so train.csv / test.csv must NOT be committed to the repository. The files written here
contain no article text and are safe to commit.

Inputs (keep them OUT of the repo):
    <data-dir>/train.csv, <data-dir>/test.csv     columns: text, label, y   (written by the Kaggle notebook)
    --audit-xlsx label_audit_18_items.xlsx        optional, your filled-in label audit
Outputs in --results-dir:
    preds_TFIDF.npy      TF-IDF + LinearSVC predictions on the test set
    test_labels.csv      columns: test_row, label, y   (no text)
    label_audit.csv      optional; columns: test_row, gold_label, check, note   (no article text)

Usage:
    python src/prepare_inputs.py --data-dir data --results-dir results --audit-xlsx label_audit_18_items.xlsx
"""
import argparse
import os

import numpy as np
import pandas as pd
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics import accuracy_score, f1_score
from sklearn.svm import LinearSVC


def tfidf_baseline(data_dir, results_dir):
    train = pd.read_csv(os.path.join(data_dir, "train.csv"))
    test = pd.read_csv(os.path.join(data_dir, "test.csv"))
    # identical configuration to the Kaggle notebook (Section 6)
    vec = TfidfVectorizer(ngram_range=(1, 2), min_df=2, max_features=200_000, sublinear_tf=True)
    clf = LinearSVC().fit(vec.fit_transform(train.text), train.y)
    pred = clf.predict(vec.transform(test.text)).astype(np.int64)
    np.save(os.path.join(results_dir, "preds_TFIDF.npy"), pred)
    test_labels = pd.DataFrame({"test_row": np.arange(len(test)), "label": test.label, "y": test.y})
    test_labels.to_csv(os.path.join(results_dir, "test_labels.csv"), index=False)
    print(f"TF-IDF + LinearSVC: accuracy={accuracy_score(test.y, pred):.4f}, "
          f"macro-F1={f1_score(test.y, pred, average='macro'):.4f}  (paper: 0.935 / 0.935)")
    print(f"Wrote preds_TFIDF.npy and test_labels.csv ({len(test)} rows, no text) to '{results_dir}'")


def audit_to_csv(xlsx_path, results_dir):
    from openpyxl import load_workbook          # only needed for this optional step
    ws = load_workbook(xlsx_path, data_only=True).active
    header_row = next(r for r in range(1, ws.max_row + 1)
                      if ws.cell(r, 1).value == "#" and ws.cell(r, 2).value == "test_row")
    cols = {str(ws.cell(header_row, c).value).strip(): c for c in range(1, ws.max_column + 1)}
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        test_row = ws.cell(r, cols["test_row"]).value
        if not isinstance(test_row, (int, float)):
            continue
        rows.append(dict(test_row=int(test_row), gold_label=ws.cell(r, cols["Gold label"]).value,
                         check=ws.cell(r, cols["Your check"]).value, note=ws.cell(r, cols["Your note"]).value))
    out = pd.DataFrame(rows)
    out.to_csv(os.path.join(results_dir, "label_audit.csv"), index=False)
    print(f"Wrote label_audit.csv ({len(out)} rows): {out.check.value_counts().to_dict()}")
    print("Check the 'note' column before publishing: it must not quote article text.")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data-dir", default="data")
    ap.add_argument("--results-dir", default="results")
    ap.add_argument("--audit-xlsx", default=None)
    a = ap.parse_args()
    os.makedirs(a.results_dir, exist_ok=True)
    tfidf_baseline(a.data_dir, a.results_dir)
    if a.audit_xlsx:
        audit_to_csv(a.audit_xlsx, a.results_dir)


if __name__ == "__main__":
    main()
